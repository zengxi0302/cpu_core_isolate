#!/usr/bin/env python3
"""Increment 2: add three more MFI deep-dive slides (T12-T14) so the
deep-dive section reads CPU x8 + MEM x6. Operates on the current deck
(after make_mfi_slides.py). Also bumps 十一大 -> 十四大 on the agenda
and divider slides."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

RED = RGBColor(0xC7, 0x00, 0x39)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xD9, 0xD9, 0xD9)
FOOT_BG = RGBColor(0xF2, 0xF2, 0xF2)
ORANGE = RGBColor(0xE8, 0x8A, 0x00)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)
CODE_HDR = RGBColor(0x2D, 0x2D, 0x44)
CODE_TITLE = RGBColor(0x6B, 0xD4, 0xA0)
CODE_CMT = RGBColor(0x7A, 0x8B, 0x99)
CODE_TXT = RGBColor(0xD4, 0xD4, 0xD4)
CHIP_BG = RGBColor(0x33, 0x33, 0x33)
YH = "Microsoft YaHei"
FOOTER_LEFT = "华为云基础设施  |  CPU / 内存故障自隔离（CFI·MFI）"

prs = Presentation("CFI_华为云_汇报材料.pptx")
BLANK = prs.slides[3].slide_layout


def box(slide, x, y, w, h, fill=None, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    return sh


def text(slide, x, y, w, h, runs, size=12, bold=False, color=BODY, font=YH,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(45720)
    tf.margin_top = tf.margin_bottom = Emu(22860)
    if isinstance(runs, str):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, {})]
        for txt, ov in para:
            r = p.add_run()
            r.text = txt
            r.font.name = ov.get("font", font)
            r.font.size = Pt(ov.get("size", size))
            r.font.bold = ov.get("bold", bold)
            r.font.color.rgb = ov.get("color", color)
            rPr = r._r.get_or_add_rPr()
            ea = rPr.find(qn("a:ea"))
            if ea is None:
                ea = rPr.makeelement(qn("a:ea"), {})
                rPr.append(ea)
            ea.set("typeface", ov.get("font", font))
    return tb


def footer(slide):
    box(slide, 0, 7.18, 13.333, 0.32, fill=FOOT_BG)
    text(slide, 0.40, 7.18, 6.6, 0.32, FOOTER_LEFT, size=9, color=GRAY)
    text(slide, 8.93, 7.18, 4.0, 0.32, "华为机密  Huawei Confidential",
         size=9, color=GRAY)


def dive_header(slide, tag, title, tnum, difficulty):
    box(slide, 0, 0, 13.333, 1.15, fill=WHITE)
    box(slide, 0.50, 0.28, 0.12, 0.55, fill=RED)
    box(slide, 0.75, 0.26, 2.0, 0.32, fill=CHIP_BG)
    text(slide, 0.75, 0.26, 2.0, 0.32, tag, size=11, bold=True, color=WHITE,
         font="Arial", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.75, 0.60, 11.0, 0.5, title, size=23, bold=True, color=DARK)
    text(slide, 11.83, 0.30, 1.0, 0.5, tnum, size=20, bold=True, color=RED,
         font="Arial")
    box(slide, 0.75, 1.12, 11.8, 0.02, fill=LIGHT)
    text(slide, 9.60, 0.62, 0.95, 0.30, "技术难度", size=10.5, bold=True,
         color=GRAY)
    for i in range(5):
        box(slide, 10.55 + 0.26 * i, 0.67, 0.20, 0.20,
            fill=RED if i < difficulty else LIGHT)


def dive_card(slide, y, h, bar_color, heading, body_lines):
    box(slide, 0.75, y, 5.85, h, fill=WHITE, line=LIGHT)
    box(slide, 0.75, y, 0.09, h, fill=bar_color)
    text(slide, 1.00, y + 0.12, 5.45, 0.36, "▎" + heading, size=12.5,
         bold=True, color=bar_color)
    text(slide, 1.00, y + 0.50, 5.45, h - 0.60, body_lines, size=11.5,
         color=BODY, line_spacing=1.05)


def code_block(slide, header, lines):
    box(slide, 6.90, 1.40, 5.65, 0.34, fill=CODE_HDR)
    text(slide, 7.05, 1.40, 5.35, 0.34, header, size=11, bold=True,
         color=CODE_TITLE, font="Consolas", anchor=MSO_ANCHOR.MIDDLE)
    box(slide, 6.90, 1.74, 5.65, 4.66, fill=CODE_BG)
    paras = []
    for ln in lines:
        color = CODE_CMT if ln.lstrip().startswith(("/*", "//", "*", "#")) \
            else CODE_TXT
        paras.append([(ln if ln else " ", {"color": color})])
    text(slide, 7.08, 1.84, 5.31, 4.46, paras, size=10.5, font="Consolas",
         line_spacing=1.0)


# ============ T12: 有界记账 ============
s = prs.slides.add_slide(BLANK)
dive_header(s, "BOUNDED STATE", "攻坚点十二：为「海量页」记账，而不是为「百个核」", "T12", 3)
dive_card(s, 1.40, 1.70, RED, "难点",
          ["CPU 域的状态对象只有 nr_cpu_ids 个，静态数组即可；",
           "内存域面对的是 TB 级宿主机的数亿个物理页。",
           "事件还来自原子上下文（tracepoint / MCE 通路），",
           "不能睡眠、不能大分配，但页处置本身必须能睡眠。"])
dive_card(s, 3.20, 1.55, ORANGE, "常规做法为何失效",
          ["• 为每页建状态 → struct page 不可改（ko 约束），",
           "  外挂全量表内存失控",
           "• 无界哈希随错误增长 → 坏 DIMM 风暴打爆内核内存",
           "• 在事件上下文直接做页迁移 → 原子上下文睡眠，死机"])
dive_card(s, 4.85, 1.55, GREEN, "我们的解法",
          ["固定容量哈希（1024 项，最坏 ~64KB）+ LRU 淘汰，",
           "在途页（PRE_ISO/POISONED）永不淘汰；OFFLINED 即出表，",
           "内核 HWPoison 标志做持久事实源；事件路径只做",
           "查表+计数（GFP_ATOMIC），soft offline 等重操作",
           "全部下沉独立 workqueue，卸载时 destroy 自动排空。"])
code_block(s, "core/mfi_core.c  ·  有界表 + 在途保护", [
    "/* 容量满: 沿 LRU 找可牺牲项 */",
    "list_for_each_entry(victim, &lru, lru) {",
    "    /* 在途页不可淘汰: 丢的是状态,",
    "     * 不是正在执行的隔离动作 */",
    "    if (victim->state == PRE_ISO ||",
    "        victim->state == POISONED)",
    "        continue;",
    "    hash_del(&victim->hash); ...",
    "}",
    "",
    "/* 隔离成功 => 状态出表, 事实交给",
    " * 内核持久标志, 表只存\"进行中\" */",
    "if (success)",
    "    mfi_page_remove(e);  /* PageHWPoison",
    "                            为准 */",
    "",
    "/* 事件路径: GFP_ATOMIC + spinlock,",
    " * 页迁移下沉 mfi_offline workqueue */",
])
footer(s)
T12 = s

# ============ T13: 事件源融合 ============
s = prs.slides.add_slide(BLANK)
dive_header(s, "EVENT FUSION", "攻坚点十三：四个事件源，一个事实——融合与归一化", "T13", 4)
dive_card(s, 1.40, 1.70, RED, "难点",
          ["同一个物理错误可能同时从四个口进来：MCE decode",
           "chain（x86）、EDAC mc_event（带 DIMM label）、GHES",
           "（arm64）、FMA（HCE3）。地址语义不一（MCi_ADDR vs",
           "EDAC 解码地址）、信息互补（有的带 label 无 PFN，",
           "有的带 PFN 无 label）、EDAC 不区分是否已消费。"])
dive_card(s, 3.20, 1.55, ORANGE, "常规做法为何失效",
          ["• 只挑一个源 → x86 丢 DIMM 定位，arm64 丢一切",
           "• 多源各自处理 → 同一错误重复计数、重复隔离",
           "• tracepoint 原型逐参数硬匹配，内核版本间漂移",
           "  → 编译期不报错，运行时探针静默错位"])
dive_card(s, 4.85, 1.55, GREEN, "我们的解法",
          ["全部归一到 mfi_mem_error{pfn,type,flags,label}：",
           "消费语义只信 MCE/SEA（AR 位），EDAC UCE 一律按",
           "异步处理交给 HWPoison 判重兜底；无地址的 mc_event",
           "只做介质记账；探针签名对照目标内核 ras_event.h",
           "校验，注册失败降级为「该源不可用」而非整体失败。"])
code_block(s, "core/mfi_dimm.c  ·  归一化与降级", [
    "/* EDAC 视角没有\"是否已消费\":",
    " * 一律按 deferred 报, 由 HWPoison",
    " * 判重避免与 MCE 路径双重处置 */",
    "case HW_EVENT_ERR_UNCORRECTED:",
    "    err.type = MFI_MEM_UCE_DEFERRED;",
    "",
    "/* 信息互补: 无地址 => 只记介质 */",
    "if (!address) {",
    "    mfi_dimm_account(label, uce);",
    "    return;",
    "}",
    "err.pfn = address >> PAGE_SHIFT;",
    "",
    "/* 源不可用 != 模块失败 */",
    "if (register_trace_mc_event(..)) {",
    "    pr_warn(\"DIMM accounting off\");",
    "    return 0;  /* MCE/GHES 路径仍在 */",
    "}",
])
footer(s)
T13 = s

# ============ T14: 可验证性 ============
s = prs.slides.add_slide(BLANK)
dive_header(s, "TESTABILITY", "攻坚点十四：如何测试一条「判决可能是 panic」的路径", "T14", 3)
dive_card(s, 1.40, 1.70, RED, "难点",
          ["甄别引擎的一半出口是 panic——跑一次错判决，测试机",
           "就没了；真实 UCE 需要 EINJ 物理机，迭代以天计；",
           "页隔离会真杀进程，随机选靶页可能误伤测试环境",
           "自身。逻辑正确性不能依赖「上机碰运气」。"])
dive_card(s, 3.20, 1.55, ORANGE, "常规做法为何失效",
          ["• 全靠物理机硬注入 → 一轮迭代一天，判决表",
           "  组合根本扫不全",
           "• mock 内核接口做单测 → mock 与真实语义漂移，",
           "  测的是 mock 不是代码"])
dive_card(s, 4.85, 1.55, GREEN, "我们的解法",
          ["四层金字塔：① 判决逻辑抽成零内核依赖纯函数",
           "（mfi_policy.h），内核与单测共享同一份实现，",
           "18 用例扫全判决表；② debugfs 注入走真实处置路径；",
           "③ ownpage 工具 mmap+mlock 自报 PFN，杀的只是",
           "牺牲进程；④ EINJ 物理机只验通路，不验逻辑。"])
code_block(s, "test/  ·  分层验证（逻辑与通路解耦）", [
    "/* L1: 纯函数单测, 内核同源同实现 */",
    "#include \"../../kernel/core/mfi_policy.h\"",
    "CHECK(decide(PG_USER,  1, 0) == RECOVER);",
    "CHECK(decide(PG_KERNEL,1, 0) == PANIC);",
    "CHECK(decide(PG_USER,  0, 0) == PANIC);",
    "/* ... 18 cases, 全绿 */",
    "",
    "# L2: 真实路径, 软件触发",
    "echo \"domain=mem pfn=$PFN type=uce_srao\"",
    "     > /sys/kernel/debug/cfi/inject",
    "",
    "# L3: 靶页自备, 误伤面=1个牺牲进程",
    "$ ./ownpage &   # 输出 pfn=0x1a2b3c",
    "$ mem_inject.sh --type uce_srar",
    "",
    "# L4: EINJ 物理机, 只验硬件通路",
    "echo 0x10 > .../einj/error_type",
])
footer(s)
T14 = s

# ============ 修订既有页 ============

def edit(slide, old, new):
    for sh in slide.shapes:
        if sh.has_text_frame and old in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        return True
    print(f"  !! not found: {old[:30]}")
    return False

edit(prs.slides[1], "关键技术与十一大攻坚点", "关键技术与十四大攻坚点")
edit(prs.slides[9], "技术深潜：十一大攻坚点", "技术深潜：十四大攻坚点")
edit(prs.slides[9], "HWPoison 页隔离复用  ·  跨域协同",
     "HWPoison 页隔离复用  ·  跨域协同  ·  有界记账  ·  多源事件融合  ·  分层可验证性")

# ============ 重排序: T12-T14 插到 T11(20) 之后 ============
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
# 当前 25 页为 ids[0..24]，新页为 ids[25..27]
order = list(range(0, 21)) + [25, 26, 27] + list(range(21, 25))
assert sorted(order) == list(range(28))
for el in ids:
    sldIdLst.remove(el)
for idx in order:
    sldIdLst.append(ids[idx])

prs.save("CFI_华为云_汇报材料.pptx")
print("saved:", len(list(prs.slides)), "slides")
