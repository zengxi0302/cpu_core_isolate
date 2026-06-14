#!/usr/bin/env python3
"""Increment 4 (v0.7): erase the self-built lockup detection / isolation
feature from the deck. The capability was removed from the kernel module —
software lockups dominate (>95%) and auto-isolating them does more harm
than good — so the deck should not advertise it any longer.

Operations:
  1. Drop slide 15 WATCHDOG (the "自建双探针" technical dive, T5) entirely.
  2. Patch slide 16 STATE MACHINE in place: remove the "Lockup ↘" arrow
     label and the "Lockup 特殊路径" card; rewrite the code-block body to
     drop the CFI_ERR_HARDLOCKUP / CFI_ERR_SOFTLOCKUP branch.
  3. Renumber the agenda + deep-dive divider: 十七大 -> 十六大, and remove
     "WATCHDOG ·" from the dive chip list on the divider.
  4. Same renumbering on subsequent T-tags is not required: T6..T17 are
     just textual chips that read fine after a count change; we leave
     them untouched to avoid disturbing layout. Only the divider chip
     list reflects feature names, not slide numbers.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

RED = RGBColor(0xC7, 0x00, 0x39)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
CODE_TXT = RGBColor(0xD4, 0xD4, 0xD4)
CODE_CMT = RGBColor(0x7A, 0x8B, 0x99)

prs = Presentation("CFI_华为云_汇报材料.pptx")


# ----- 1. Drop WATCHDOG slide (index 14, 1-indexed 15) -----
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
victim = ids[14]
sldIdLst.remove(victim)
# also drop the relationship so the slide is fully gone
rId = victim.get(qn("r:id"))
prs.part.drop_rel(rId)
print(f"removed slide id={victim.get('id')} rId={rId} (WATCHDOG)")


# ----- helpers for in-place edits -----
def shape_text(sh):
    return sh.text_frame.text if sh.has_text_frame else ""


def drop_shape(sh):
    sp = sh._element
    sp.getparent().remove(sp)


def rewrite_shape_text(sh, new_lines, color=None):
    """Replace a shape's text frame with new paragraph list.
    new_lines: list of str or list of [(text, opts_dict), ...]
    """
    tf = sh.text_frame
    # nuke all existing paragraphs except the first
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    first.clear()
    for i, line in enumerate(new_lines):
        p = first if i == 0 else tf.add_paragraph()
        if isinstance(line, str):
            line = [(line, {})]
        for txt, ov in line:
            r = p.add_run()
            r.text = txt
            r.font.name = ov.get("font", "Consolas")
            r.font.size = Pt(ov.get("size", 10.5))
            r.font.bold = ov.get("bold", False)
            r.font.color.rgb = ov.get("color", color or CODE_TXT)


# ----- 2. Patch STATE MACHINE slide (was index 15, now 14 after delete) -----
# After drop, indices shift. Resolve by content.
state_slide = None
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame and "五态状态机" in sh.text_frame.text:
            state_slide = s
            break
    if state_slide:
        break

assert state_slide is not None, "STATE MACHINE slide not found"

# drop "Lockup ↘" arrow label + the "Lockup 特殊路径" card heading + body
removed = 0
shapes_to_drop = []
for sh in state_slide.shapes:
    t = shape_text(sh)
    if not t.strip():
        continue
    if t.strip() == "Lockup ↘":
        shapes_to_drop.append(sh); removed += 1; continue
    if "Lockup 特殊路径" in t and "▎" in t:
        shapes_to_drop.append(sh); removed += 1; continue
    if t.startswith("硬锁 / 软锁事件跳过阈值累计"):
        shapes_to_drop.append(sh); removed += 1; continue

for sh in shapes_to_drop:
    drop_shape(sh)
print(f"STATE MACHINE: dropped {removed} lockup shapes")

# rewrite the code-block lines to remove the lockup branch
for sh in state_slide.shapes:
    t = shape_text(sh)
    if "CFI_ERR_HARDLOCKUP" in t or ("cfi_transition" in t and "ISOLATING" in t):
        new = [
            [("/* 累计阈值 → 状态转移 */",
              {"color": CODE_CMT, "size": 11, "font": "Consolas"})],
            [("if (sev >= UCE) {",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("    ci->uce_count++;",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("    if (ci->uce_count >= cfi_uce_threshold)",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("        cfi_transition(ci, cpu,",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("                       CFI_STATE_ISOLATING);",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("} else if (ci->ce_count >= cfi_ce_threshold) {",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("    cfi_transition(ci, cpu,",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("                   CFI_STATE_DEGRADED);",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("}",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [(" ",
              {"color": CODE_TXT, "size": 11, "font": "Consolas"})],
            [("/* ONLINE ──CE≥10──> DEGRADED ──UCE≥1──> ISOLATING */",
              {"color": CODE_CMT, "size": 11, "font": "Consolas"})],
        ]
        rewrite_shape_text(sh, new)
        print("STATE MACHINE: code block rewritten")
        break


# ----- 3. Renumber agenda + deep-dive divider -----
def edit(slide, old, new):
    for sh in slide.shapes:
        if sh.has_text_frame and old in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        return True
    print(f"  !! edit not found: {old[:40]}")
    return False


# Slides after the WATCHDOG drop: agenda is slide 2 (idx 1), divider was
# slide 10 (idx 9), now still idx 9 because everything before it untouched.
edit(prs.slides[1], "十七大", "十六大")
edit(prs.slides[9], "十七大", "十六大")
# Strip WATCHDOG from the dive chip list on the divider
edit(prs.slides[9],
     "TIMING  ·  KALLSYMS  ·  FMA  ·  CONCURRENCY  ·  WATCHDOG  ·  STATE MACHINE",
     "TIMING  ·  KALLSYMS  ·  FMA  ·  CONCURRENCY  ·  STATE MACHINE")
# (some decks list chips differently — try common variants)
for sl in [prs.slides[9]]:
    for sh in sl.shapes:
        if sh.has_text_frame and "WATCHDOG" in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if "WATCHDOG" in r.text:
                        r.text = r.text.replace("  ·  WATCHDOG", "")
                        r.text = r.text.replace("WATCHDOG  ·  ", "")
                        r.text = r.text.replace("WATCHDOG", "")


prs.save("CFI_华为云_汇报材料.pptx")
print(f"saved: {len(prs.slides._sldIdLst)} slides")
